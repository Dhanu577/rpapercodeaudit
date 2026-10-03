from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path
from random import Random
from typing import Any

from openpyxl import load_workbook

SEED = 42
PINNED_COMMIT = "76c5f8523716804dbe0a9500b4b7e216c6af225c"
ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPO = ROOT / "examples" / "deseq2" / "repo"
DEFAULT_PAPER = ROOT / "examples" / "deseq2" / "DESeq.txt"
DEFAULT_TABLE = ROOT / "examples" / "deseq2" / "final" / "deseq2_final_table.xlsx"


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text)).strip()


def git(repo: Path, *args: str) -> str:
    return subprocess.run(["git", "-C", str(repo), *args], check=True, text=True, capture_output=True).stdout.strip()


def git_show(repo: Path, commit: str, path: str) -> str | None:
    proc = subprocess.run(["git", "-C", str(repo), "show", f"{commit}:{path}"], text=True, capture_output=True)
    return proc.stdout if proc.returncode == 0 else None


def load_records(table: Path) -> list[dict[str, Any]]:
    ws = load_workbook(table, read_only=True, data_only=True)["Final Table"]
    headers = list(next(ws.iter_rows(values_only=True)))
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        match = re.fullmatch(r"\s*(\d+)\s*-\s*(\d+)\s*", str(row.get("Code lines", "")))
        if not match or not row.get("Code file") or not row.get("Code excerpt") or not row.get("Paper sentence"):
            continue
        rows.append({
            "record_id": f"row-{int(row['Original row'])}",
            "source_record_id": f"row-{int(row['Original row'])}",
            "paper_sentence": str(row["Paper sentence"]),
            "file": str(row["Code file"]),
            "line_start": int(match.group(1)),
            "line_end": int(match.group(2)),
            "excerpt": str(row["Code excerpt"]),
            "justification": str(row.get("Justification") or ""),
        })
    return rows


def base(record: dict[str, Any], *, case_id: str, tier: str, mutation_type: str, params: dict[str, Any], sentence: str | None = None, commit: str = PINNED_COMMIT, file: str | None = None, start: int | None = None, end: int | None = None, excerpt: str | None = None, justification: str | None = None, expected_valid: bool = False, expected_failing_stage: str | None = None) -> dict[str, Any]:
    return {
        "case_id": case_id, "source_record_id": record["source_record_id"], "tier": tier,
        "mutation_type": mutation_type, "variant_params": params,
        "paper_sentence": sentence if sentence is not None else record["paper_sentence"],
        "repo": "thelovelab/DESeq2", "commit": commit,
        "file": file if file is not None else record["file"],
        "line_start": start if start is not None else record["line_start"],
        "line_end": end if end is not None else record["line_end"],
        "excerpt": excerpt if excerpt is not None else record["excerpt"],
        "justification": justification if justification is not None else record["justification"],
        "expected_valid": expected_valid, "expected_failing_stage": expected_failing_stage, "seed": SEED,
    }


def altered_sentence(sentence: str, variant: int) -> str:
    if re.search(r"\d+(?:\.\d+)?", sentence):
        return re.sub(r"\d+(?:\.\d+)?", lambda m: str(int(float(m.group(0))) + variant + 1), sentence, count=1)
    words = sentence.split()
    if len(words) > 3:
        words[-2] = f"not_{words[-2].strip('.,;:')}{variant}"
    return " ".join(words)


def text_files(repo: Path, commit: str) -> list[tuple[str, str]]:
    paths = git(repo, "ls-tree", "-r", "--name-only", commit).splitlines()
    allowed = {".r", ".R", ".cpp", ".cc", ".c", ".h", ".hpp", ".rd", ".md", ".rmd", ".tex", ".txt"}
    out = []
    for path in paths:
        if not any(path.startswith(prefix + "/") for prefix in ("R", "src", "man", "vignettes", "tests")):
            continue
        if Path(path).suffix not in allowed:
            continue
        content = git_show(repo, commit, path)
        if content is not None:
            out.append((path, content))
    return out


def choose_wrong_commit(repo: Path, record: dict[str, Any]) -> str | None:
    for commit in git(repo, "rev-list", "--all", "--max-count=80").splitlines():
        if commit == PINNED_COMMIT:
            continue
        content = git_show(repo, commit, record["file"])
        if content is None:
            continue
        lines = content.splitlines(keepends=True)
        if record["line_end"] <= len(lines):
            actual = "".join(lines[record["line_start"] - 1:record["line_end"]])
            if norm(actual) != norm(record["excerpt"]):
                return commit
    return None


def build_cases(repo: Path, paper_path: Path, table_path: Path) -> list[dict[str, Any]]:
    records = load_records(table_path)
    paper = norm(paper_path.read_text(encoding="utf-8"))
    rng = Random(SEED)
    cases: list[dict[str, Any]] = []
    for record in records:
        rid = record["source_record_id"]
        cases.append(base(record, case_id=f"{rid}-A0", tier="A", mutation_type="A0_valid", params={}, expected_valid=True))
        for variant in range(3):
            cases.append(base(record, case_id=f"{rid}-A1-{variant+1}", tier="A", mutation_type="A1_altered_sentence", params={"variant": variant + 1}, sentence=altered_sentence(record["paper_sentence"], variant), expected_failing_stage="sentence_match"))
        other = records[(records.index(record) + 1) % len(records)]
        for variant in range(3):
            left = record["paper_sentence"][: max(1, len(record["paper_sentence"]) // 2)]
            right = other["paper_sentence"][max(1, len(other["paper_sentence"]) // 2):]
            sentence = norm(left + " " + right)
            if sentence in paper:
                sentence = f"{left} {right} This fabricated clause is absent {variant}."
            cases.append(base(record, case_id=f"{rid}-A2-{variant+1}", tier="A", mutation_type="A2_nonexistent_sentence", params={"spliced_with": other["source_record_id"], "variant": variant + 1}, sentence=sentence, expected_failing_stage="sentence_match"))
        wrong_commit = choose_wrong_commit(repo, record)
        if wrong_commit:
            cases.append(base(record, case_id=f"{rid}-A3", tier="A", mutation_type="A3_wrong_commit", params={"commit_differs": True}, commit=wrong_commit, expected_failing_stage="repo_commit_valid"))
        candidates = [(path, content) for path, content in text_files(repo, PINNED_COMMIT) if path != record["file"] and len(content.splitlines()) >= record["line_end"]]
        if candidates:
            wrong_file, _ = candidates[rng.randrange(len(candidates))]
            cases.append(base(record, case_id=f"{rid}-A4", tier="A", mutation_type="A4_wrong_file", params={"excerpt_source_file": record["file"]}, file=wrong_file, expected_failing_stage="excerpt_matches"))
        pinned_text = git_show(repo, PINNED_COMMIT, record["file"]) or ""
        pinned_line_count = len(pinned_text.splitlines())
        plus3_failure = "line_range_valid" if record["line_end"] + 3 > pinned_line_count else "excerpt_matches"
        cases.extend([
            base(record, case_id=f"{rid}-A5-shift-minus3", tier="A", mutation_type="A5_invalid_line_range", params={"variant": "shift_-3"}, start=max(1, record["line_start"] - 3), end=max(1, record["line_end"] - 3), expected_failing_stage="excerpt_matches"),
            base(record, case_id=f"{rid}-A5-shift-plus3", tier="A", mutation_type="A5_invalid_line_range", params={"variant": "shift_+3"}, start=record["line_start"] + 3, end=record["line_end"] + 3, expected_failing_stage=plus3_failure),
            base(record, case_id=f"{rid}-A5-start-after-end", tier="A", mutation_type="A5_invalid_line_range", params={"variant": "start_gt_end"}, start=record["line_end"], end=record["line_start"], expected_failing_stage="line_range_valid"),
            base(record, case_id=f"{rid}-A5-zero-start", tier="A", mutation_type="A5_invalid_line_range", params={"variant": "zero_start"}, start=0, end=record["line_end"], expected_failing_stage="line_range_valid"),
            base(record, case_id=f"{rid}-A5-missing-file", tier="A", mutation_type="A5_invalid_line_range", params={"variant": "nonexistent_file"}, file="R/__benchmark_missing__.R", expected_failing_stage="file_exists"),
        ])
        # Tier B deliberately uses authentic repository text but tests semantic boundaries.
        cases.append(base(record, case_id=f"{rid}-B1", tier="B", mutation_type="B1_irrelevant_authentic_excerpt", params={"semantic_scope": "irrelevant_same_repo"}, file=other["file"], start=other["line_start"], end=other["line_end"], excerpt=other["excerpt"], expected_valid=True))
        cases.append(base(record, case_id=f"{rid}-B2", tier="B", mutation_type="B2_misleading_justification", params={"semantic_scope": "overstated_justification"}, justification="This excerpt proves every downstream implementation detail and all semantic consequences.", expected_valid=True))
        cases.append(base(record, case_id=f"{rid}-B3", tier="B", mutation_type="B3_contradictory_authentic_excerpt", params={"semantic_scope": "authentic_but_contradictory"}, file=other["file"], start=other["line_start"], end=other["line_end"], excerpt=other["excerpt"], expected_valid=True))
        cases.append(base(record, case_id=f"{rid}-B4", tier="B", mutation_type="B4_compound_claim", params={"semantic_scope": "compound_claim_interpretation", "second_claim": other["source_record_id"]}, justification=f"The authentic sentence is interpreted as a compound claim; the cited excerpt supports only one component. {record['justification']}", expected_valid=True))
    return cases


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", type=Path, default=DEFAULT_REPO)
    parser.add_argument("--paper", type=Path, default=DEFAULT_PAPER)
    parser.add_argument("--table", type=Path, default=DEFAULT_TABLE)
    parser.add_argument("--output", type=Path, default=ROOT / "benchmark" / "cases.jsonl")
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    cases = build_cases(args.repo, args.paper, args.table)
    text = "".join(json.dumps(case, sort_keys=True, ensure_ascii=False) + "\n" for case in cases)
    args.output.write_text(text, encoding="utf-8")
    checksum = hashlib.sha256(text.encode()).hexdigest()
    args.output.with_name("cases.sha256").write_text(f"{checksum}  {args.output.name}\n", encoding="utf-8")
    print(json.dumps({"cases": len(cases), "checksum": checksum, "seed": SEED, "source_records": len(load_records(args.table))}, indent=2))

if __name__ == "__main__":
    main()
